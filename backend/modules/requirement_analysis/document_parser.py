# Extracts requirement text from supported uploaded documents.
#
# Supported formats:
# - TXT
# - PDF containing selectable/extractable text
# - DOCX
#
# Scanned/image-only PDFs are intentionally not supported.

from io import BytesIO
from pathlib import Path

from docx import Document
from pypdf import PdfReader


SUPPORTED_EXTENSIONS = {
    ".txt",
    ".pdf",
    ".docx",
}


class DocumentParsingError(Exception):
    """
    Raised when an uploaded document cannot be parsed.
    """

    pass


def get_file_extension(
    filename: str
) -> str:

    return (
        Path(filename)
        .suffix
        .lower()
    )


def validate_file_type(
    filename: str
):

    extension = get_file_extension(
        filename
    )

    if extension not in SUPPORTED_EXTENSIONS:

        supported = ", ".join(
            sorted(
                SUPPORTED_EXTENSIONS
            )
        )

        raise DocumentParsingError(
            f"Unsupported file type: "
            f"{extension or 'unknown'}. "
            f"Supported formats: {supported}"
        )


def extract_txt_text(
    file_bytes: bytes
) -> str:

    try:

        return file_bytes.decode(
            "utf-8"
        )

    except UnicodeDecodeError:

        try:

            return file_bytes.decode(
                "utf-8-sig"
            )

        except UnicodeDecodeError as exc:

            raise DocumentParsingError(
                "Unable to decode the TXT file "
                "as UTF-8 text."
            ) from exc


def extract_pdf_text(
    file_bytes: bytes
) -> str:

    try:

        reader = PdfReader(
            BytesIO(
                file_bytes
            )
        )

    except Exception as exc:

        raise DocumentParsingError(
            "Unable to open the PDF document."
        ) from exc


    extracted_pages = []


    for page in reader.pages:

        try:

            page_text = (
                page.extract_text()
                or ""
            )

        except Exception:

            page_text = ""


        page_text = (
            page_text.strip()
        )

        if page_text:

            extracted_pages.append(
                page_text
            )


    text = "\n\n".join(
        extracted_pages
    )


    if not text.strip():

        raise DocumentParsingError(
            "No extractable text was found in "
            "the PDF. The file may be scanned "
            "or image-only."
        )


    return text


def extract_docx_text(
    file_bytes: bytes
) -> str:

    try:

        document = Document(
            BytesIO(
                file_bytes
            )
        )

    except Exception as exc:

        raise DocumentParsingError(
            "Unable to open the DOCX document."
        ) from exc


    content = []


    # Extract normal paragraphs.

    for paragraph in document.paragraphs:

        text = (
            paragraph.text
            .strip()
        )

        if text:

            content.append(
                text
            )


    # Extract text stored inside tables.

    for table in document.tables:

        for row in table.rows:

            row_values = []

            for cell in row.cells:

                cell_text = (
                    cell.text
                    .strip()
                )

                if cell_text:

                    row_values.append(
                        cell_text
                    )


            if row_values:

                content.append(
                    " | ".join(
                        row_values
                    )
                )


    text = "\n".join(
        content
    )


    if not text.strip():

        raise DocumentParsingError(
            "No extractable text was found "
            "in the DOCX document."
        )


    return text


def clean_extracted_text(
    text: str
) -> str:

    lines = []


    for line in text.splitlines():

        cleaned = (
            line
            .strip()
        )

        if cleaned:

            lines.append(
                cleaned
            )


    return "\n".join(
        lines
    )


def extract_document_text(
    filename: str,
    file_bytes: bytes,
) -> str:

    if not filename:

        raise DocumentParsingError(
            "Uploaded file has no filename."
        )


    if not file_bytes:

        raise DocumentParsingError(
            "Uploaded document is empty."
        )


    validate_file_type(
        filename
    )


    extension = get_file_extension(
        filename
    )


    if extension == ".txt":

        text = extract_txt_text(
            file_bytes
        )


    elif extension == ".pdf":

        text = extract_pdf_text(
            file_bytes
        )


    elif extension == ".docx":

        text = extract_docx_text(
            file_bytes
        )


    else:

        raise DocumentParsingError(
            "Unsupported document format."
        )


    cleaned_text = (
        clean_extracted_text(
            text
        )
    )


    if len(cleaned_text) < 10:

        raise DocumentParsingError(
            "The uploaded document does not "
            "contain enough readable text "
            "for requirement analysis."
        )


    return cleaned_text