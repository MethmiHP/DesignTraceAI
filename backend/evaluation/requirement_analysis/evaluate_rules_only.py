# Evaluates the deterministic requirement-quality rules against the
# human-reviewed gold dataset. The script measures rule-only performance
# for ambiguity and incompleteness detection using accuracy, precision,
# recall, F1-score, and confusion-matrix counts.

from pathlib import Path

from openpyxl import load_workbook

from modules.requirement_analysis.rules import (
    run_quality_rules
)


# FILE PATH

GOLD_FILE = Path(
    "datasets/requirement_analysis/human_gold/"
    "requirement_analysis_gold.xlsx"
)


# HELPERS
def to_bool(value):

    if isinstance(value, bool):
        return value

    if value is None:
        return False

    return str(value).strip().lower() in {
        "true",
        "1",
        "yes"
    }


def safe_divide(
    numerator,
    denominator
):

    if denominator == 0:
        return 0.0

    return numerator / denominator


# LOAD GOLD DATASET

def load_gold_records():

    if not GOLD_FILE.exists():

        raise FileNotFoundError(
            f"Gold file not found: "
            f"{GOLD_FILE}"
        )

    workbook = load_workbook(
        GOLD_FILE,
        data_only=True
    )

    sheet = workbook[
        workbook.sheetnames[0]
    ]

    rows = list(
        sheet.iter_rows(
            values_only=True
        )
    )

    if not rows:

        raise ValueError(
            "Gold workbook is empty."
        )

    headers = [

        str(value).strip()
        if value is not None
        else ""

        for value in rows[0]
    ]

    records = []

    for row in rows[1:]:

        record = dict(
            zip(
                headers,
                row
            )
        )

        if not record.get(
            "requirement_id"
        ):
            continue

        records.append({

            "project_id":
                record.get(
                    "project_id"
                ),

            "requirement_id":
                record.get(
                    "requirement_id"
                ),

            "statement":
                record.get(
                    "statement"
                ),

            "gold_ambiguous":
                to_bool(
                    record.get(
                        "final_ambiguous"
                    )
                ),

            "gold_incomplete":
                to_bool(
                    record.get(
                        "final_incomplete"
                    )
                ),
        })

    return records


# BINARY METRICS

def binary_metrics(
    y_true,
    y_pred
):

    tp = 0
    tn = 0
    fp = 0
    fn = 0

    for truth, pred in zip(
        y_true,
        y_pred
    ):

        if truth is True and pred is True:

            tp += 1

        elif truth is False and pred is False:

            tn += 1

        elif truth is False and pred is True:

            fp += 1

        elif truth is True and pred is False:

            fn += 1

    total = (
        tp
        + tn
        + fp
        + fn
    )

    accuracy = safe_divide(
        tp + tn,
        total
    )

    precision = safe_divide(
        tp,
        tp + fp
    )

    recall = safe_divide(
        tp,
        tp + fn
    )

    f1 = safe_divide(
        2 * precision * recall,
        precision + recall
    )

    return {

        "accuracy":
            accuracy,

        "precision":
            precision,

        "recall":
            recall,

        "f1":
            f1,

        "tp":
            tp,

        "tn":
            tn,

        "fp":
            fp,

        "fn":
            fn,

        "total":
            total,
    }


# PRINT METRICS

def print_metrics(
    title,
    metrics
):

    print(
        "\n======================================="
    )

    print(
        title
    )

    print(
        "======================================="
    )

    print(
        f"Total     : "
        f"{metrics['total']}"
    )

    print(
        f"Accuracy  : "
        f"{metrics['accuracy']:.4f}"
    )

    print(
        f"Precision : "
        f"{metrics['precision']:.4f}"
    )

    print(
        f"Recall    : "
        f"{metrics['recall']:.4f}"
    )

    print(
        f"F1        : "
        f"{metrics['f1']:.4f}"
    )

    print(
        f"TP={metrics['tp']} "
        f"TN={metrics['tn']} "
        f"FP={metrics['fp']} "
        f"FN={metrics['fn']}"
    )


# MAIN

def main():

    print(
        "\n======================================="
    )

    print(
        "RULE-ONLY QUALITY EVALUATION"
    )

    print(
        "======================================="
    )


    # STEP 1 - LOAD GOLD

    records = load_gold_records()

    print(
        f"Gold records loaded: "
        f"{len(records)}"
    )


    # STEP 2 - RUN RULES

    ambiguous_true = []

    ambiguous_pred = []

    incomplete_true = []

    incomplete_pred = []


    print(
        "\nRunning deterministic rules..."
    )


    for index, record in enumerate(
        records,
        start=1
    ):

        result = run_quality_rules(
            record[
                "statement"
            ]
        )


        gold_ambiguous = (
            record[
                "gold_ambiguous"
            ]
        )

        pred_ambiguous = (
            result.get(
                "ambiguous",
                False
            )
        )


        gold_incomplete = (
            record[
                "gold_incomplete"
            ]
        )

        pred_incomplete = (
            result.get(
                "incomplete",
                False
            )
        )


        ambiguous_true.append(
            gold_ambiguous
        )

        ambiguous_pred.append(
            pred_ambiguous
        )


        incomplete_true.append(
            gold_incomplete
        )

        incomplete_pred.append(
            pred_incomplete
        )


        print(
            f"[{index}/{len(records)}] "
            f"{record['project_id']} "
            f"{record['requirement_id']} "
            f"| Amb Gold={gold_ambiguous} "
            f"Pred={pred_ambiguous} "
            f"| Inc Gold={gold_incomplete} "
            f"Pred={pred_incomplete}"
        )


    # AMBIGUITY METRICS
    ambiguous_metrics = (
        binary_metrics(
            ambiguous_true,
            ambiguous_pred
        )
    )

    print_metrics(

        "A. RULE-ONLY AMBIGUITY DETECTION",

        ambiguous_metrics
    )


    # INCOMPLETENESS METRICS

    incomplete_metrics = (
        binary_metrics(
            incomplete_true,
            incomplete_pred
        )
    )

    print_metrics(

        "B. RULE-ONLY INCOMPLETENESS DETECTION",

        incomplete_metrics
    )


    # FINAL SUMMARY

    print(
        "\n======================================="
    )

    print(
        "FINAL RULE-ONLY SUMMARY"
    )

    print(
        "======================================="
    )

    print(
        "\nAmbiguity:"
    )

    print(
        f"Accuracy  = "
        f"{ambiguous_metrics['accuracy']:.4f}"
    )

    print(
        f"Precision = "
        f"{ambiguous_metrics['precision']:.4f}"
    )

    print(
        f"Recall    = "
        f"{ambiguous_metrics['recall']:.4f}"
    )

    print(
        f"F1        = "
        f"{ambiguous_metrics['f1']:.4f}"
    )


    print(
        "\nIncompleteness:"
    )

    print(
        f"Accuracy  = "
        f"{incomplete_metrics['accuracy']:.4f}"
    )

    print(
        f"Precision = "
        f"{incomplete_metrics['precision']:.4f}"
    )

    print(
        f"Recall    = "
        f"{incomplete_metrics['recall']:.4f}"
    )

    print(
        f"F1        = "
        f"{incomplete_metrics['f1']:.4f}"
    )


# ENTRY POINT

if __name__ == "__main__":

    main()