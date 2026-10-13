# Loads the human-reviewed Requirement Analysis gold dataset from Excel,
# converts the final annotation fields into a standardized Python format,
# normalizes boolean labels, and prints a summary for verification before
# using the gold data in model and pipeline evaluation.

import json
from pathlib import Path

from openpyxl import load_workbook


GOLD_FILE = Path(
    "datasets/requirement_analysis/human_gold/"
    "requirement_analysis_gold.xlsx"
)


def to_bool(value):
    """
    Convert spreadsheet values into Python booleans.
    """

    if isinstance(value, bool):
        return value

    if value is None:
        return False

    value = str(value).strip().lower()

    return value in {
        "true",
        "1",
        "yes"
    }


def main():

    if not GOLD_FILE.exists():
        raise FileNotFoundError(
            f"Gold file not found: {GOLD_FILE}"
        )

    workbook = load_workbook(
        GOLD_FILE,
        data_only=True
    )

    sheet = workbook[
        workbook.sheetnames[0]
    ]

    rows = list(
        sheet.iter_rows(
            values_only=True
        )
    )

    if not rows:
        raise ValueError(
            "Gold workbook is empty."
        )

    headers = [
        str(value).strip()
        if value is not None
        else ""
        for value in rows[0]
    ]

    records = []

    for row in rows[1:]:

        record = dict(
            zip(
                headers,
                row
            )
        )

        if not record.get(
            "requirement_id"
        ):
            continue

        gold_record = {
            "project_id":
                record.get(
                    "project_id"
                ),

            "requirement_id":
                record.get(
                    "requirement_id"
                ),

            "statement":
                record.get(
                    "statement"
                ),

            "requirement_type":
                record.get(
                    "final_requirement_type"
                ),

            "quality_attribute":
                record.get(
                    "final_quality_attribute"
                ),

            "asr":
                to_bool(
                    record.get(
                        "final_asr"
                    )
                ),

            "ambiguous":
                to_bool(
                    record.get(
                        "final_ambiguous"
                    )
                ),

            "incomplete":
                to_bool(
                    record.get(
                        "final_incomplete"
                    )
                ),

            "conflicting":
                to_bool(
                    record.get(
                        "final_conflicting"
                    )
                ),

            "potentially_missing":
                to_bool(
                    record.get(
                        "final_potentially_missing"
                    )
                ),
        }

        records.append(
            gold_record
        )

    print(
        "\n==================================="
    )

    print(
        "REQUIREMENT ANALYSIS GOLD DATASET"
    )

    print(
        "==================================="
    )

    print(
        f"Total records: {len(records)}"
    )

    project_ids = sorted(
        {
            record["project_id"]
            for record in records
        }
    )

    print(
        f"Projects: {len(project_ids)}"
    )

    print(
        f"Project IDs: {project_ids}"
    )

    print(
        "\nFirst 5 gold records:\n"
    )

    for record in records[:5]:

        print(
            json.dumps(
                record,
                indent=2
            )
        )

        print(
            "-" * 50
        )


if __name__ == "__main__":
    main()