# Evaluates the selected trained Requirement Analysis classifiers directly
# against the human-reviewed gold dataset. It measures standalone FR/NFR
# classification, standalone NFR quality-attribute classification, and the
# complete FR/NFR-to-NFR-category prediction chain using multiclass metrics.

from pathlib import Path

from openpyxl import load_workbook

from modules.requirement_analysis.classifier_service import (
    get_requirement_classifier
)


# FILE PATH

GOLD_FILE = Path(
    "datasets/requirement_analysis/human_gold/"
    "requirement_analysis_gold.xlsx"
)


# HELPERS

def safe_divide(
    numerator,
    denominator
):

    if denominator == 0:
        return 0.0

    return numerator / denominator


# LOAD GOLD DATASET

def load_gold_records():

    if not GOLD_FILE.exists():

        raise FileNotFoundError(
            f"Gold file not found: "
            f"{GOLD_FILE}"
        )

    workbook = load_workbook(
        GOLD_FILE,
        data_only=True
    )

    sheet = workbook[
        workbook.sheetnames[0]
    ]

    rows = list(
        sheet.iter_rows(
            values_only=True
        )
    )

    if not rows:

        raise ValueError(
            "Gold workbook is empty."
        )

    headers = [

        str(value).strip()
        if value is not None
        else ""

        for value in rows[0]
    ]

    records = []

    for row in rows[1:]:

        record = dict(
            zip(
                headers,
                row
            )
        )

        if not record.get(
            "requirement_id"
        ):
            continue

        records.append({

            "project_id":
                record.get(
                    "project_id"
                ),

            "requirement_id":
                record.get(
                    "requirement_id"
                ),

            "statement":
                record.get(
                    "statement"
                ),

            "gold_requirement_type":
                record.get(
                    "final_requirement_type"
                ),

            "gold_quality_attribute":
                record.get(
                    "final_quality_attribute"
                ),
        })

    return records


# MULTICLASS METRICS

def multiclass_metrics(
    y_true,
    y_pred
):

    labels = sorted(
        set(y_true)
        | set(y_pred)
    )

    total = len(
        y_true
    )

    correct = sum(

        1

        for truth, pred
        in zip(
            y_true,
            y_pred
        )

        if truth == pred
    )

    accuracy = safe_divide(
        correct,
        total
    )

    per_class = {}

    macro_f1_values = []

    weighted_f1_sum = 0.0

    total_support = 0

    for label in labels:

        tp = sum(

            1

            for truth, pred
            in zip(
                y_true,
                y_pred
            )

            if (
                truth == label
                and
                pred == label
            )
        )

        fp = sum(

            1

            for truth, pred
            in zip(
                y_true,
                y_pred
            )

            if (
                truth != label
                and
                pred == label
            )
        )

        fn = sum(

            1

            for truth, pred
            in zip(
                y_true,
                y_pred
            )

            if (
                truth == label
                and
                pred != label
            )
        )

        support = sum(

            1

            for truth in y_true

            if truth == label
        )

        precision = safe_divide(
            tp,
            tp + fp
        )

        recall = safe_divide(
            tp,
            tp + fn
        )

        f1 = safe_divide(

            2
            * precision
            * recall,

            precision
            + recall
        )

        per_class[
            label
        ] = {

            "precision":
                precision,

            "recall":
                recall,

            "f1":
                f1,

            "support":
                support,
        }

        # Only include classes that actually
        # exist in the gold dataset when
        # calculating macro F1.
        if support > 0:

            macro_f1_values.append(
                f1
            )

        weighted_f1_sum += (
            f1
            * support
        )

        total_support += (
            support
        )

    macro_f1 = safe_divide(

        sum(
            macro_f1_values
        ),

        len(
            macro_f1_values
        )
    )

    weighted_f1 = safe_divide(

        weighted_f1_sum,

        total_support
    )

    return {

        "accuracy":
            accuracy,

        "macro_f1":
            macro_f1,

        "weighted_f1":
            weighted_f1,

        "per_class":
            per_class,

        "total":
            total,

        "correct":
            correct,
    }


# PRINT MULTICLASS RESULTS

def print_multiclass_results(
    title,
    metrics
):

    print(
        "\n======================================="
    )

    print(
        title
    )

    print(
        "======================================="
    )

    print(
        f"Total        : "
        f"{metrics['total']}"
    )

    print(
        f"Correct      : "
        f"{metrics['correct']}"
    )

    print(
        f"Accuracy     : "
        f"{metrics['accuracy']:.4f}"
    )

    print(
        f"Macro F1     : "
        f"{metrics['macro_f1']:.4f}"
    )

    print(
        f"Weighted F1  : "
        f"{metrics['weighted_f1']:.4f}"
    )

    print(
        "\nPer-class results:"
    )

    for label, values in (
        metrics[
            "per_class"
        ].items()
    ):

        # Skip classes that appear only
        # in predictions but not in gold.
        if values[
            "support"
        ] == 0:

            continue

        print(
            f"{label}: "
            f"Precision="
            f"{values['precision']:.4f}, "
            f"Recall="
            f"{values['recall']:.4f}, "
            f"F1="
            f"{values['f1']:.4f}, "
            f"Support="
            f"{values['support']}"
        )


# MAIN

def main():

    print(
        "\n======================================="
    )

    print(
        "CLASSIFIER-ONLY GOLD EVALUATION"
    )

    print(
        "======================================="
    )


    # STEP 1 - LOAD GOLD DATA

    records = (
        load_gold_records()
    )

    print(
        f"Gold records loaded: "
        f"{len(records)}"
    )


    # STEP 2 - LOAD CLASSIFIERS
    print(
        "\nLoading trained classifiers..."
    )

    classifier = (
        get_requirement_classifier()
    )

    print(
        "Classifiers loaded."
    )


    # A. STANDALONE FR / NFR EVALUATION

    type_true = []

    type_pred = []

    print(
        "\nRunning standalone "
        "FR/NFR evaluation..."
    )

    for index, record in enumerate(
        records,
        start=1
    ):

        result = (
            classifier
            .classify_requirement_type(
                record[
                    "statement"
                ]
            )
        )

        gold_type = (
            record[
                "gold_requirement_type"
            ]
        )

        predicted_type = (
            result[
                "requirement_type"
            ]
        )

        type_true.append(
            gold_type
        )

        type_pred.append(
            predicted_type
        )

        print(
            f"[{index}/{len(records)}] "
            f"{record['project_id']} "
            f"{record['requirement_id']} "
            f"| Gold={gold_type} "
            f"| Pred={predicted_type}"
        )

    type_metrics = (
        multiclass_metrics(
            type_true,
            type_pred
        )
    )

    print_multiclass_results(

        "A. STANDALONE FR / NFR CLASSIFIER",

        type_metrics
    )


    # B. STANDALONE NFR CATEGORY EVALUATION

    gold_nfr_records = [

        record

        for record in records

        if (
            record[
                "gold_requirement_type"
            ]
            == "NFR"
        )
    ]

    standalone_nfr_true = []

    standalone_nfr_pred = []

    print(
        "\nRunning standalone "
        "NFR-category evaluation..."
    )

    for index, record in enumerate(
        gold_nfr_records,
        start=1
    ):

        result = (
            classifier
            .classify_nfr_category(
                record[
                    "statement"
                ]
            )
        )

        gold_category = (
            record[
                "gold_quality_attribute"
            ]
            or ""
        )

        predicted_category = (
            result[
                "quality_attribute"
            ]
            or ""
        )

        standalone_nfr_true.append(
            gold_category
        )

        standalone_nfr_pred.append(
            predicted_category
        )

        print(
            f"[{index}/"
            f"{len(gold_nfr_records)}] "
            f"{record['project_id']} "
            f"{record['requirement_id']} "
            f"| Gold={gold_category} "
            f"| Pred={predicted_category}"
        )

    standalone_nfr_metrics = (
        multiclass_metrics(
            standalone_nfr_true,
            standalone_nfr_pred
        )
    )

    print_multiclass_results(

        "B. STANDALONE NFR CATEGORY CLASSIFIER",

        standalone_nfr_metrics
    )


    # C. END-TO-END CLASSIFICATION CHAIN

    chain_true = []

    chain_pred = []

    print(
        "\nRunning end-to-end "
        "classification-chain evaluation..."
    )

    for index, record in enumerate(
        gold_nfr_records,
        start=1
    ):

        result = (
            classifier.classify(
                record[
                    "statement"
                ]
            )
        )

        gold_category = (
            record[
                "gold_quality_attribute"
            ]
            or ""
        )

        # If FR/NFR predicts FR,
        # the NFR category stage is never called.
        #
        # Record it explicitly as TYPE_ERROR.
        if (
            result[
                "requirement_type"
            ]
            != "NFR"
        ):

            predicted_category = (
                "TYPE_ERROR"
            )

        else:

            predicted_category = (
                result[
                    "quality_attribute"
                ]
                or "NO_CATEGORY"
            )

        chain_true.append(
            gold_category
        )

        chain_pred.append(
            predicted_category
        )

        print(
            f"[{index}/"
            f"{len(gold_nfr_records)}] "
            f"{record['project_id']} "
            f"{record['requirement_id']} "
            f"| Gold={gold_category} "
            f"| Pred={predicted_category}"
        )

    chain_metrics = (
        multiclass_metrics(
            chain_true,
            chain_pred
        )
    )

    print_multiclass_results(

        "C. END-TO-END NFR CLASSIFICATION CHAIN",

        chain_metrics
    )


    # FINAL SUMMARY

    print(
        "\n======================================="
    )

    print(
        "FINAL CLASSIFIER EVALUATION SUMMARY"
    )

    print(
        "======================================="
    )

    print(
        "\nFR/NFR standalone:"
    )

    print(
        f"Accuracy    = "
        f"{type_metrics['accuracy']:.4f}"
    )

    print(
        f"Macro F1    = "
        f"{type_metrics['macro_f1']:.4f}"
    )

    print(
        f"Weighted F1 = "
        f"{type_metrics['weighted_f1']:.4f}"
    )

    print(
        "\nNFR category standalone:"
    )

    print(
        f"Accuracy    = "
        f"{standalone_nfr_metrics['accuracy']:.4f}"
    )

    print(
        f"Macro F1    = "
        f"{standalone_nfr_metrics['macro_f1']:.4f}"
    )

    print(
        f"Weighted F1 = "
        f"{standalone_nfr_metrics['weighted_f1']:.4f}"
    )

    print(
        "\nNFR category end-to-end chain:"
    )

    print(
        f"Accuracy    = "
        f"{chain_metrics['accuracy']:.4f}"
    )

    print(
        f"Macro F1    = "
        f"{chain_metrics['macro_f1']:.4f}"
    )

    print(
        f"Weighted F1 = "
        f"{chain_metrics['weighted_f1']:.4f}"
    )


# ENTRY POINT

if __name__ == "__main__":

    main()